"""
Module xuất báo cáo phân tích sang định dạng Excel (.xlsx) đa trang
Học phần: DATA ANALYSIS FOR BUSINESS ENVIRONMENT (71DAEE10012)
"""

import sys
from pathlib import Path
from typing import Dict, Optional, Tuple

import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Đường dẫn mặc định trong project
DEFAULT_ROOT = Path(__file__).resolve().parents[3]
TABLES_DIR = DEFAULT_ROOT / "outputs" / "tables"
REPORTS_DIR = DEFAULT_ROOT / "outputs" / "reports"


def export_excel_report(output_file: Optional[Path] = None) -> Path:
    """
    Xuất toàn bộ các bảng kết quả phân tích sang một file Excel đa trang (Multi-sheet Workbook)
    với thiết kế chuẩn Dashboard chuyên nghiệp:
    - Sheet 1: Executive_Summary (KPIs tổng quan & Tóm tắt phân khúc)
    - Sheet 2: Customer_Segments (Chân dung 3 hạng thẻ hội viên)
    - Sheet 3: Monthly_Revenue (Doanh thu 13 tháng: Gross, Cancelled, Net)
    - Sheet 4: Country_Revenue (Thị phần theo từng quốc gia)
    - Sheet 5: Top_Products (Sản phẩm đóng góp doanh số cao nhất)
    - Sheet 6: Statistical_Tests (Kiểm định Mann-Whitney U & Chi-Square)
    - Sheet 7: KMeans_Metrics (Chỉ số đánh giá K=2 đến K=7)
    - Sheet 8: Invoice_Distribution (Phân phối giá trị đơn hàng)
    """
    if output_file is None:
        REPORTS_DIR.mkdir(parents=True, exist_ok=True)
        output_file = REPORTS_DIR / "Online_Retail_Analysis_Report.xlsx"
    else:
        output_file = Path(output_file)
        output_file.parent.mkdir(parents=True, exist_ok=True)

    # 1. Khởi tạo Workbook
    wb = openpyxl.Workbook()
    default_sheet = wb.active

    # Bảng màu & Định dạng chuẩn
    NAVY_HEADER = "1F4E79"
    SOFT_BLUE = "D9E1F2"
    LIGHT_GRAY = "F2F2F2"
    BORDER_COLOR = "D9D9D9"

    font_title = Font(name="Segoe UI", size=14, bold=True, color="1F4E79")
    font_subtitle = Font(name="Segoe UI", size=10, italic=True, color="595959")
    font_section = Font(name="Segoe UI", size=11, bold=True, color="1F4E79")
    font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Segoe UI", size=10, color="000000")
    font_bold = Font(name="Segoe UI", size=10, bold=True, color="000000")

    fill_header = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    fill_zebra = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type="solid")

    thin_side = Side(border_style="thin", color=BORDER_COLOR)
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    # -------------------------------------------------------------
    # SHEET 1: EXECUTIVE SUMMARY
    # -------------------------------------------------------------
    ws1 = wb.create_sheet(title="Executive_Summary")
    ws1.views.sheetView[0].showGridLines = True

    # Tiêu đề
    ws1["B2"] = "BÁO CÁO PHÂN TÍCH DOANH THU & PHÂN CỤM KHÁCH HÀNG"
    ws1["B2"].font = font_title
    ws1["B3"] = "Học phần: DATA ANALYSIS FOR BUSINESS ENVIRONMENT (71DAEE10012) – Sem1/2026-2027"
    ws1["B3"].font = font_subtitle

    # Khung KPI Tổng Quan
    kpis = [
        ("Tổng Doanh Thu Thuần", 8291748.56, "£#,##0.00"),
        ("Tổng Số Hóa Đơn Thành Công", 19960, "#,##0"),
        ("Số Khách Hàng Định Danh", 4339, "#,##0"),
        ("Giá Trị Đơn Trung Vị (Median)", 303.30, "£#,##0.00"),
        ("Tỷ Trọng Doanh Thu UK", 0.8401, "0.00%"),
        ("Tỷ Lệ Hủy Đơn Tại Đức", 0.2420, "0.00%"),
    ]

    ws1["B5"] = "I. CHỈ SỐ VẬN HÀNH & KINH DOANH THEN CHỐT"
    ws1["B5"].font = font_section

    ws1["B6"] = "Chỉ Số Đo Lường"
    ws1["C6"] = "Giá Trị Đạt Được"
    for col in ["B", "C"]:
        ws1[f"{col}6"].font = font_header
        ws1[f"{col}6"].fill = fill_header
        ws1[f"{col}6"].alignment = Alignment(horizontal="center", vertical="center")

    curr_row = 7
    for name, val, fmt in kpis:
        c1 = ws1[f"B{curr_row}"]
        c2 = ws1[f"C{curr_row}"]
        c1.value = name
        c2.value = val
        c1.font = font_data
        c2.font = font_bold
        c2.number_format = fmt
        c1.border = cell_border
        c2.border = cell_border
        c2.alignment = Alignment(horizontal="right")
        if curr_row % 2 == 1:
            c1.fill = fill_zebra
            c2.fill = fill_zebra
        curr_row += 1

    # Bảng Tóm Tắt Hạng Thẻ Ngay Sheet 1
    curr_row += 2
    ws1[f"B{curr_row}"] = "II. TÓM TẮT HỆ THỐNG HẠNG THẺ HỘI VIÊN (K-MEANS CLUSTERING)"
    ws1[f"B{curr_row}"].font = font_section
    curr_row += 1

    seg_summary_file = TABLES_DIR / "customer_segments_summary.csv"
    if seg_summary_file.exists():
        df_seg = pd.read_csv(seg_summary_file)
        headers = ["Hạng Thẻ Hội Viên", "Cụm", "Số Khách", "Tỷ Lệ Khách (%)", "R (ngày)", "F (đơn)", "M (£)", "Tổng Doanh Thu (£)", "Tỷ Trọng DT (%)"]
        start_col = 2  # Cột B

        for c_idx, h in enumerate(headers, start=start_col):
            col_letter = get_column_letter(c_idx)
            cell = ws1[f"{col_letter}{curr_row}"]
            cell.value = h
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = Alignment(horizontal="center", vertical="center")

        curr_row += 1
        for _, r in df_seg.iterrows():
            row_data = [
                r["MembershipTier"],
                int(r["Cluster"]),
                int(r["CustomerCount"]),
                r["Customer_Pct"] / 100.0,
                r["Recency_Mean"],
                r["Frequency_Mean"],
                r["Monetary_Mean"],
                r["TotalMonetary"],
                r["Revenue_Share_Pct"] / 100.0,
            ]
            for c_idx, val in enumerate(row_data, start=start_col):
                col_letter = get_column_letter(c_idx)
                cell = ws1[f"{col_letter}{curr_row}"]
                cell.value = val
                cell.font = font_data
                cell.border = cell_border
                # Format
                if c_idx in [2]:
                    cell.alignment = Alignment(horizontal="left")
                elif c_idx in [3, 4]:
                    cell.alignment = Alignment(horizontal="center")
                    if c_idx == 4:
                        cell.number_format = "#,##0"
                elif c_idx in [5, 10]:
                    cell.alignment = Alignment(horizontal="right")
                    cell.number_format = "0.00%"
                elif c_idx in [6, 7]:
                    cell.alignment = Alignment(horizontal="right")
                    cell.number_format = "#,##0.0"
                elif c_idx in [8, 9]:
                    cell.alignment = Alignment(horizontal="right")
                    cell.number_format = "£#,##0.00"
            curr_row += 1

    # Tự động chỉnh độ rộng cột Sheet 1
    for col in ws1.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # -------------------------------------------------------------
    # HÀM TRỢ GIÚP THÊM SHEET TỪ CSV CÓ SẴN
    # -------------------------------------------------------------
    def _add_csv_sheet(sheet_title: str, csv_path: Path, col_formats: Dict[int, Tuple[str, str]]):
        if not csv_path.exists():
            return
        df = pd.read_csv(csv_path)
        ws = wb.create_sheet(title=sheet_title)
        ws.views.sheetView[0].showGridLines = True

        # Header Title
        ws["B2"] = sheet_title.replace("_", " ").upper()
        ws["B2"].font = font_title
        ws["B3"] = f"Dữ liệu nguồn trích xuất từ: outputs/tables/{csv_path.name}"
        ws["B3"].font = font_subtitle

        start_r = 5
        start_c = 2

        # Header Columns
        for c_idx, col_name in enumerate(df.columns, start=start_c):
            col_letter = get_column_letter(c_idx)
            cell = ws[f"{col_letter}{start_r}"]
            cell.value = col_name
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Data Rows
        for r_idx, row in df.iterrows():
            current_r = start_r + 1 + r_idx
            for c_idx, col_name in enumerate(df.columns, start=start_c):
                col_letter = get_column_letter(c_idx)
                cell = ws[f"{col_letter}{current_r}"]
                val = row[col_name]

                cell.value = val
                cell.font = font_data
                cell.border = cell_border

                if r_idx % 2 == 1:
                    cell.fill = fill_zebra

                if c_idx - start_c in col_formats:
                    fmt_align, fmt_num = col_formats[c_idx - start_c]
                    cell.alignment = Alignment(horizontal=fmt_align)
                    if fmt_num:
                        cell.number_format = fmt_num
                else:
                    if isinstance(val, (int, float)):
                        cell.alignment = Alignment(horizontal="right")
                    else:
                        cell.alignment = Alignment(horizontal="left")

        # Auto width
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Thêm các sheets dữ liệu chi tiết
    _add_csv_sheet(
        "Customer_Segments",
        TABLES_DIR / "customer_segments_summary.csv",
        {
            0: ("left", None),
            1: ("center", "#,##0"),
            2: ("right", "#,##0"),
            3: ("right", "0.00%"),
            4: ("right", "#,##0.0"),
            5: ("right", "#,##0.0"),
            6: ("right", "#,##0.0"),
            7: ("right", "#,##0.0"),
            8: ("right", "£#,##0.00"),
            9: ("right", "£#,##0.00"),
            10: ("right", "£#,##0.00"),
            11: ("right", "0.00%"),
        }
    )

    _add_csv_sheet(
        "Monthly_Revenue",
        TABLES_DIR / "monthly_revenue.csv",
        {
            0: ("center", None),
            1: ("right", "£#,##0.00"),
            2: ("right", "£#,##0.00"),
            3: ("right", "£#,##0.00"),
            4: ("right", "#,##0"),
            5: ("right", "#,##0"),
        }
    )

    _add_csv_sheet(
        "Country_Revenue",
        TABLES_DIR / "country_revenue.csv",
        {
            0: ("left", None),
            1: ("right", "£#,##0.00"),
            2: ("right", "#,##0"),
            3: ("right", "#,##0"),
            4: ("right", "0.00%"),
        }
    )

    _add_csv_sheet(
        "Top_Products",
        TABLES_DIR / "top_products_revenue.csv",
        {
            0: ("center", None),
            1: ("left", None),
            2: ("right", "£#,##0.00"),
            3: ("right", "#,##0"),
            4: ("right", "#,##0"),
        }
    )

    _add_csv_sheet(
        "Statistical_Tests",
        TABLES_DIR / "statistical_tests_summary.csv",
        {
            0: ("left", None),
            1: ("left", None),
            2: ("left", None),
            3: ("center", None),
            4: ("center", None),
            5: ("center", None),
            6: ("left", None),
        }
    )

    _add_csv_sheet(
        "KMeans_Metrics",
        TABLES_DIR / "kmeans_evaluation_metrics.csv",
        {
            0: ("center", "#,##0"),
            1: ("right", "#,##0.00"),
            2: ("right", "0.0000"),
        }
    )

    _add_csv_sheet(
        "Invoice_Distribution",
        TABLES_DIR / "invoice_value_summary.csv",
        {
            0: ("left", None),
            1: ("right", "#,##0.00"),
        }
    )

    if default_sheet in wb.worksheets:
        wb.remove(default_sheet)

    wb.save(output_file)
    print(f"[EXCEL EXPORT] Đã xuất thành công file Excel báo cáo: {output_file}")
    return output_file
